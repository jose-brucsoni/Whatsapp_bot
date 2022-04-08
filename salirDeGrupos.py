from selenium import webdriver
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support import expected_conditions as ec
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.common.by import By
import time
from buscarGrupo import buscarGrupo
from datetime import datetime
from openpyxl  import Workbook
from openpyxl import load_workbook
from remove import remove
from salirGrupo import salirGrupo


#Se Carga el archivo Excel
excel = load_workbook('modulo2.xlsx')
#Selecciona la hoja dentro del excel
ws = excel['modulo2']
#Selecciona el rango de la tabla
cell_range = ws['A2':'T500']

lst=[]


chrome = webdriver.Edge(executable_path='./msedgedriver')

wait = WebDriverWait(chrome,600)

chrome.get('https://web.whatsapp.com/')

contador = -1

for i in cell_range:

    contador += 1
    groupName = cell_range[contador][12].value
    
    if(groupName == "" or groupName == None):
            break

    if(cell_range[contador][17].value == "" or cell_range[contador][17].value == None):

        
            #Aqui empieza a Buscar el grupo
            x_nameFieldGroup = '//*[@id="side"]/div[1]/div/label/div/div[2]'
            nameFieldGroup = wait.until(ec.presence_of_element_located((By.XPATH,x_nameFieldGroup)))
            nameFieldGroup.clear()
            nameFieldGroup.send_keys(groupName)
            time.sleep(6)
            nameFieldGroup.send_keys(Keys.ENTER)
            time.sleep(2)

            try:
                x_textoResultadoDeBusqueda = chrome.find_element_by_xpath('//*[@id="pane-side"]/div/div')
                textoResultadoDeBusqueda = str(x_textoResultadoDeBusqueda.get_attribute("textContent"))

                

                if(textoResultadoDeBusqueda == "No se encontró ningún chat, contacto ni mensaje"):

                    print("NO SE ENCONTRO NINGUN GRUPO CON EL NOMBRE DE => ",groupName)

                else:

                    #Click a Boton De 3 puntos
                    x_threePointsButton = '//*[@id="main"]/header/div[3]/div/div[2]/div/div'
                    buttonthreePoints = wait.until(ec.presence_of_element_located((By.XPATH,x_threePointsButton)))
                    buttonthreePoints.click()

                    time.sleep(2)

                    #Click en Info Grupo
                    x_InfoGrupoButton = '//*[@id="app"]/div[1]/span[4]/div/ul/div/div/li[1]'
                    buttonInfoGrupo = wait.until(ec.presence_of_element_located((By.XPATH,x_InfoGrupoButton)))
                    buttonInfoGrupo.click()


                    time.sleep(2)

                    panel = chrome.find_element_by_xpath('//*[@id="app"]/div[1]/div[1]/div[2]/div[3]/span/div[1]/span/div[1]/div')
                    a = chrome.execute_script("return arguments[0].scrollHeight;",panel)
                    for i in range (0,a):
                        try:
                            s=chrome.find_element_by_xpath('//*[@id="app"]/div[1]/div[1]/div[2]/div[3]/span/div[1]/span/div[1]/div/section/div[6]/div[2]/div/div/div[8]/div/div/div[2]/div[1]/div[2]')

                            if (ec.presence_of_element_located((By.XPATH,s))):
                                print ("se encontro a Diana")
                                break
                        except:
                            chrome.execute_script("arguments[0].scroll(0,arguments[1]);",panel,i*10)

                    time.sleep(3)



                    try:

                        try:
                            #extrae si es "administrador" desde la ventana "Participantes", el segundo contacto mostrado
                            x_VerificarSiEsAdmin = chrome.find_element_by_xpath('//*[@id="app"]/div[1]/div[1]/div[2]/div[3]/span/div[1]/span/div[1]/div/section/div[6]/div[2]/div/div/div[8]/div/div/div[2]/div[1]/div[2]')
                            verificarSiEsAdmin = str(x_VerificarSiEsAdmin.get_attribute("textContent"))

                            #Extrae el nombre del contacto del segundo contacto mostrado en lista
                            x_nombreDecontacto_1 = chrome.find_element_by_xpath('//*[@id="app"]/div[1]/div[1]/div[2]/div[3]/span/div[1]/span/div[1]/div/section/div[6]/div[2]/div[3]/div/div[1]/div/div/div[2]/div[1]/div')
                            nombreDecontacto_1 = str(x_nombreDecontacto_1.get_attribute("textContent"))

                            #Extrae el nombre del contacto del tercer contacto mostrado en lista
                            x_nombreDecontacto_2 = chrome.find_element_by_xpath('//*[@id="app"]/div[1]/div[1]/div[2]/div[3]/span/div[1]/span/div[1]/div/section/div[6]/div[2]/div[3]/div/div[7]/div/div/div[2]/div[1]/div')
                            nombreDecontacto_2 = str(x_nombreDecontacto_2.get_attribute("textContent"))


                            if(nombreDecontacto_1[0] != "+" and nombreDecontacto_2[0] == "+"):

                                if(verificarSiEsAdmin == "Admin. del grupo"):

                                    print("El Grupo: ", groupName, ", ya tiene administrador")

                                    salirGrupo(wait,groupName,chrome,cell_range, contador,excel)
                                    print("SE ELIMINO EL GRUPO ===>", groupName)
                                    

                            elif(nombreDecontacto_1[0] != "+" and nombreDecontacto_2[0] != "+"):

                                print("----------------El Grupo: ", groupName, ", contiene 2 docentes, PERO EXISTE ADMINISTRADOR-------------")
                                salirGrupo(wait,groupName,chrome,cell_range, contador,excel)

                            else:
                                print("----------------El Grupo: ", groupName, ", NO TIENE NINGUN NUMERO REGISTRADO, Pero tiene administradores")
                                
                        except:
                            #No encontro administrador
                            # #Salir de la ventana participantes
                            # x_salirDelaVentanaParticipantes = '//*[@id="app"]/div[1]/span[2]/div[1]/div/div/div/div/div/div/header/div/div[1]/button/span'
                            # salirDelaVentanaParticipantes = wait.until(ec.presence_of_element_located((By.XPATH,x_salirDelaVentanaParticipantes)))
                            # salirDelaVentanaParticipantes.click()
                            print("El GRUPO: '", groupName,"' NO TIENE ADMINISTRADOR DESIGNADO---")
                            cell_range[contador][17].value = "NO TIENE ADMINISTRADOR DESIGNADO"
                            excel.save('modulo2.xlsx')


                    except:
                        
                        print("El GRUPO: '", groupName,"' NO POSEE ADMINISTRADOR")
                        cell_range[contador][17].value = "NO TIENE ADMINISTRADOR DESIGNADO"

                        #Confirma el guardado de los datos en Excel
                        excel.save('modulo2.xlsx')


                        print("GRUPO CON 2 PARTICIPANTES(CElULAR 1 Y 2)")
                    
                    
            except:

                #Click a Boton De 3 puntos
                x_threePointsButton = '//*[@id="main"]/header/div[3]/div/div[2]/div/div'
                buttonthreePoints = wait.until(ec.presence_of_element_located((By.XPATH,x_threePointsButton)))
                buttonthreePoints.click()
                time.sleep(2)

                #Click en Info Grupo
                x_InfoGrupoButton = '//*[@id="app"]/div[1]/span[4]/div/ul/div/div/li[1]'
                buttonInfoGrupo = wait.until(ec.presence_of_element_located((By.XPATH,x_InfoGrupoButton)))
                buttonInfoGrupo.click()

                time.sleep(2)

                x_scrollDeInfor = '//*[@id="app"]/div[1]/div[1]/div[2]/div[3]/span/div[1]/span/div[1]/div'
                scrollDeInfor = wait.until(ec.presence_of_element_located((By.XPATH,x_scrollDeInfor)))
                scrollDeInfor.execute_script("window.scrollTo(0, 1000)")

                time.sleep(2)
                try:

                    try:
                        #extrae si es "administrador" desde la ventana "Participantes", el segundo contacto mostrado
                        x_VerificarSiEsAdmin = chrome.find_element_by_xpath('//*[@id="app"]/div[1]/div[1]/div[2]/div[3]/span/div[1]/span/div[1]/div/section/div[6]/div[2]/div/div/div[8]/div/div/div[2]/div[1]/div[2]')
                        verificarSiEsAdmin = str(x_VerificarSiEsAdmin.get_attribute("textContent"))

                        #Extrae el nombre del contacto del segundo contacto mostrado en lista
                        x_nombreDecontacto_1 = chrome.find_element_by_xpath('//*[@id="app"]/div[1]/div[1]/div[2]/div[3]/span/div[1]/span/div[1]/div/section/div[6]/div[2]/div[3]/div/div[1]/div/div/div[2]/div[1]/div')
                        nombreDecontacto_1 = str(x_nombreDecontacto_1.get_attribute("textContent"))

                        #Extrae el nombre del contacto del tercer contacto mostrado en lista
                        x_nombreDecontacto_2 = chrome.find_element_by_xpath('//*[@id="app"]/div[1]/div[1]/div[2]/div[3]/span/div[1]/span/div[1]/div/section/div[6]/div[2]/div[3]/div/div[7]/div/div/div[2]/div[1]/div')
                        nombreDecontacto_2 = str(x_nombreDecontacto_2.get_attribute("textContent"))

                        print(verificarSiEsAdmin,nombreDecontacto_1[0],nombreDecontacto_2[0],"-----------------------------------------")

                        if(nombreDecontacto_1[0] != "+" and nombreDecontacto_2[0] == "+"):

                            if(verificarSiEsAdmin == "Admin. del grupo"):

                                print("El Grupo: ", groupName, ", ya tiene administrador")
                                

                                salirGrupo(wait,groupName,chrome,cell_range, contador,excel)

                            


                        elif(nombreDecontacto_1[0] != "+" and nombreDecontacto_2[0] != "+"):

                            print("----------------El Grupo: ", groupName, ", contiene 2 docentes, PERO EXISTE ADMINISTRADOR-------------")

                            salirGrupo(wait,groupName,chrome,cell_range, contador,excel)

                        else:
                            print("----------------El Grupo: ", groupName, ", NO TIENE NINGUN NUMERO REGISTRADO, Pero tiene administradores")
                            




                    except:
                        #No encontro administrador
                        
                        print("El GRUPO: ", groupName," NO TIENE ADMINISTRADOR DESIGNADO")
                        cell_range[contador][17].value = "NO TIENE ADMINISTRADOR DESIGNADO"
                        #Confirma el guardado de los datos en Excel
                        excel.save('modulo2.xlsx')

                except:

                    print("GRUPO CON 2 PARTICIPANTES(CElULAR 1 Y 2)")
                    cell_range[contador][17].value = "NO TIENE ADMINISTRADOR DESIGNADO"
                    #Confirma el guardado de los datos en Excel
                    excel.save('modulo2.xlsx')

                